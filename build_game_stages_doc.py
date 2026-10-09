from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).parent
OUTPUT = ROOT / "Wolfish_Revenge_Stages.docx"
IMAGES = ROOT / "assets" / "cutscenes" / "forest"
CAMP_IMAGES = ROOT / "assets" / "cutscenes" / "camp"
MIKHALYCH_IMAGES = ROOT / "assets" / "cutscenes" / "mikhaylych_boss"
BTR_IMAGES = ROOT / "assets" / "cutscenes" / "btr_boss"
MILITARY_UNIT_IMAGES = ROOT / "assets" / "cutscenes" / "military_unit"


def set_run_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_caption(doc, number, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(14)
    if number:
        label = p.add_run(f"Кадр {number}. ")
        set_run_font(label, size=9, bold=True, color=(80, 80, 80))
    body = p.add_run(text)
    set_run_font(body, size=9, color=(80, 80, 80))


def add_image(doc, filename, number, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(IMAGES / filename), width=Inches(6.45))
    add_caption(doc, number, caption)


def add_storyboard_map(doc):
    doc.add_heading("Схема кампании", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(ROOT / "assets" / "storyboard" / "user_game_flow_map.png"), width=Inches(6.45))
    add_caption(doc, "", "Основной маршрут и необязательные боковые локации.")


def add_cutscene_strip(doc):
    table = doc.add_table(rows=2, cols=3)
    table.autofit = False
    image_specs = [
        ("forest_cutscene_01_abduction.png", "1. Похищение девушки"),
        ("forest_cutscene_02_wolves_pov.png", "2. Волки вокруг героя"),
        ("forest_cutscene_03_transformation.png", "3. Первое превращение"),
    ]
    for cell, (filename, _) in zip(table.rows[0].cells, image_specs):
        cell.width = Inches(2.12)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(IMAGES / filename), width=Inches(2.05))
    for cell, (_, caption) in zip(table.rows[1].cells, image_specs):
        cell.width = Inches(2.12)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(caption), size=8.5, color=(80, 80, 80))


def add_camp_cutscene_strip(doc):
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    image_specs = [
        ("camp_cutscene_01_stealth_pounce.png", "1. Скрытное нападение на военного"),
        ("camp_cutscene_02_overview.png", "2. Общий вид лагеря под дождём"),
    ]
    for cell, (filename, _) in zip(table.rows[0].cells, image_specs):
        cell.width = Inches(3.18)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(CAMP_IMAGES / filename), width=Inches(3.1))
    for cell, (_, caption) in zip(table.rows[1].cells, image_specs):
        cell.width = Inches(3.18)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(caption), size=8.5, color=(80, 80, 80))


def add_mikhaylych_cutscene_strip(doc):
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    image_specs = [
        ("mikhaylych_boss_01_intro.png", "1. Появление Михалыча"),
        ("mikhaylych_boss_02_arena.png", "2. Поле сражения"),
    ]
    for cell, (filename, _) in zip(table.rows[0].cells, image_specs):
        cell.width = Inches(3.18)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(MIKHALYCH_IMAGES / filename), width=Inches(3.1))
    for cell, (_, caption) in zip(table.rows[1].cells, image_specs):
        cell.width = Inches(3.18)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(caption), size=8.5, color=(80, 80, 80))


def add_btr_intro(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(BTR_IMAGES / "btr_boss_01_intro.png"), width=Inches(5.0))
    add_caption(doc, 1, "Комиксный стопкадр появления БТР.")


def add_military_unit_storyboard(doc):
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    image_specs = [
        ("military_unit_01_facade.png", "1. Фасад и главный вход"),
        ("military_unit_02_interior.png", "2. Укрытие за внутренними перегородками"),
    ]
    for cell, (filename, _) in zip(table.rows[0].cells, image_specs):
        cell.width = Inches(3.18)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(MILITARY_UNIT_IMAGES / filename), width=Inches(3.1))
    for cell, (_, caption) in zip(table.rows[1].cells, image_specs):
        cell.width = Inches(3.18)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(caption), size=8.5, color=(80, 80, 80))


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)

for style_name, size in [("Title", 25), ("Heading 1", 17), ("Heading 2", 12)]:
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name != "Heading 2" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)

subtitle_style = styles.add_style("Subtitle Gray", WD_STYLE_TYPE.PARAGRAPH)
subtitle_style.font.name = "Aptos"
subtitle_style.font.size = Pt(11)
subtitle_style.font.color.rgb = RGBColor(89, 89, 89)

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(6)
title.add_run("Wolfish Revenge Этапы игры")

subtitle = doc.add_paragraph(style="Subtitle Gray")
subtitle.paragraph_format.space_after = Pt(22)
subtitle.add_run("Рабочий документ с краткими описаниями локаций и визуальными кадрами")

add_storyboard_map(doc)
doc.add_page_break()

doc.add_heading("1 Лес", level=1)

summary = doc.add_table(rows=3, cols=2)
summary.style = "Table Grid"
summary.autofit = False
summary.columns[0].width = Inches(1.4)
summary.columns[1].width = Inches(5.0)
rows = [
    ("Статус", "Обязательная стартовая локация"),
    ("Назначение", "Ввод в сюжет и первое знакомство с превращением в оборотня"),
    ("Задача", "Освоить превращение в волка и пробежать через лес до лагеря"),
]
for row, (label, value) in zip(summary.rows, rows):
    row.cells[0].width = Inches(1.4)
    row.cells[1].width = Inches(5.0)
    set_cell_shading(row.cells[0], "E7EBEF")
    p = row.cells[0].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(label), size=9.5, bold=True)
    p = row.cells[1].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(value), size=9.5)

doc.add_paragraph()
doc.add_heading("Завязка", level=2)
intro = (
    "Игрок начинает в тёмном лесу рядом со своей палаткой и потухшим костром. "
    "Военные увозят его девушку, а героя оглушают. Очнувшись, он от первого лица "
    "видит окруживших его волков. Затем, выйдя в лунный свет, герой впервые превращается в оборотня."
)
p = doc.add_paragraph(intro)
p.paragraph_format.space_after = Pt(12)

doc.add_heading("Кадры катсцены", level=2)
add_cutscene_strip(doc)

doc.add_heading("2 Лагерь", level=1)

summary = doc.add_table(rows=3, cols=2)
summary.style = "Table Grid"
summary.autofit = False
rows = [
    ("Статус", "Обязательная локация"),
    ("Назначение", "Первое боевое проникновение на военный объект"),
    ("Задача", "Пройти через лагерь, используя скрытность или открытый бой; после первого устранения врага получить пистолет"),
]
for row, (label, value) in zip(summary.rows, rows):
    row.cells[0].width = Inches(1.4)
    row.cells[1].width = Inches(5.0)
    set_cell_shading(row.cells[0], "E7EBEF")
    p = row.cells[0].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(label), size=9.5, bold=True)
    p = row.cells[1].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(value), size=9.5)

doc.add_paragraph()
doc.add_heading("Описание", level=2)
p = doc.add_paragraph(
    "Ночной военный палаточный лагерь в лесу. С этого этапа начинается дождь. "
    "На территории расположены около десяти крупных армейских палаток, складские ящики, "
    "песчаные мешки и две вышки с прожекторами."
)
p.paragraph_format.space_after = Pt(8)

doc.add_heading("Прохождение", level=2)
p = doc.add_paragraph(
    "Локацию можно пройти скрытно: использовать тень, шум дождя, палатки и складские ящики "
    "как укрытия, а также бесшумно устранять одиночных патрульных. Открытый бой остаётся допустимым вариантом."
)
p.paragraph_format.space_after = Pt(12)

doc.add_heading("Награда и темп", level=2)
p = doc.add_paragraph(
    "После первого убийства игрок получает пистолет и может применять его до конца этапа. "
    "Лагерь короткий и почти обучающий: он знакомит с выбором между стелсом, ближней атакой и огнестрельным боем."
)
p.paragraph_format.space_after = Pt(12)

doc.add_heading("Кадры катсцены", level=2)
add_camp_cutscene_strip(doc)

doc.add_heading("3 Ангары", level=1)

summary = doc.add_table(rows=3, cols=2)
summary.style = "Table Grid"
summary.autofit = False
rows = [
    ("Статус", "Обязательная локация"),
    ("Назначение", "Более серьёзный военный объект и подготовка к первому боссу"),
    ("Задача", "Пройти военные ангары, найти первый бронежилет и подготовиться к бою с боссом"),
]
for row, (label, value) in zip(summary.rows, rows):
    row.cells[0].width = Inches(1.4)
    row.cells[1].width = Inches(5.0)
    set_cell_shading(row.cells[0], "E7EBEF")
    p = row.cells[0].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(label), size=9.5, bold=True)
    p = row.cells[1].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(value), size=9.5)

doc.add_paragraph()
doc.add_heading("Описание", level=2)
p = doc.add_paragraph(
    "Военные ангары — уже более крупный и укреплённый объект, чем палаточный лагерь. "
    "В локации есть серьёзные сооружения, вышки и инфраструктура базы. Здесь игрок находит первый бронежилет."
)
p.paragraph_format.space_after = Pt(8)

doc.add_heading("Отдельный складской ангар", level=2)
p = doc.add_paragraph(
    "В зоне ангаров расположен отдельный складской ангар, где игрок может найти автомат Калашникова."
)
p.paragraph_format.space_after = Pt(8)

doc.add_heading("Переход к боссу", level=2)
p = doc.add_paragraph(
    "В финале локации начинается плавный переход к босс-файту. Перед ним игрока быстро знакомят "
    "с механикой вырывания дверей автомобиля: дверь используется как переносное укрытие, позволяющее "
    "подбираться к противнику с дробовиком."
)
p.paragraph_format.space_after = Pt(12)

doc.add_heading("4 Босс Михалыч", level=1)

summary = doc.add_table(rows=3, cols=2)
summary.style = "Table Grid"
summary.autofit = False
rows = [
    ("Статус", "Обязательный босс-файт"),
    ("Босс", "Михалыч — крупный мужчина в тельняшке и ватнике, вооружённый двустволкой"),
    ("Задача", "Сблизиться с боссом под огнём и победить его стрельбой или наброском"),
]
for row, (label, value) in zip(summary.rows, rows):
    row.cells[0].width = Inches(1.4)
    row.cells[1].width = Inches(5.0)
    set_cell_shading(row.cells[0], "E7EBEF")
    p = row.cells[0].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(label), size=9.5, bold=True)
    p = row.cells[1].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(value), size=9.5)

doc.add_paragraph()
doc.add_heading("Появление босса", level=2)
p = doc.add_paragraph(
    "Игрок выходит на открытую площадку. Из военного жилого прицепа на колёсах с ноги вылетает Михалыч: "
    "он держит двустволку наготове, отпивает водку и замечает героя. Появление подаётся короткой катсценой "
    "в нарочито клишированном полукомиксном стиле."
)
p.paragraph_format.space_after = Pt(8)

doc.add_heading("Арена и бой", level=2)
p = doc.add_paragraph(
    "Арена в основном открытая. Между игроком и боссом размещены автомобили, ящики, барьеры и другие препятствия, "
    "которые позволяют постепенно сокращать дистанцию перебежками. Несколько машин можно перевернуть и использовать "
    "как укрытия. Ранее освоенная автомобильная дверь также помогает двигаться под огнём двустволки."
)
p.paragraph_format.space_after = Pt(8)

doc.add_heading("Условие победы", level=2)
p = doc.add_paragraph(
    "Игроку нужно добраться до Михалыча, не получив слишком много выстрелов. После сближения босса можно "
    "застрелить или атаковать наброском."
)
p.paragraph_format.space_after = Pt(12)

doc.add_heading("Кадры катсцены", level=2)
add_mikhaylych_cutscene_strip(doc)

doc.add_heading("5 Военный городок", level=1)

summary = doc.add_table(rows=3, cols=2)
summary.style = "Table Grid"
summary.autofit = False
rows = [
    ("Статус", "Обязательная локация"),
    ("Назначение", "Освоение вертикального перемещения и прохождения через здания"),
    ("Задача", "Преодолеть низкую застройку, выбирая путь по крышам, через окна или по земле"),
]
for row, (label, value) in zip(summary.rows, rows):
    row.cells[0].width = Inches(1.4)
    row.cells[1].width = Inches(5.0)
    set_cell_shading(row.cells[0], "E7EBEF")
    p = row.cells[0].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(label), size=9.5, bold=True)
    p = row.cells[1].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(value), size=9.5)

doc.add_paragraph()
doc.add_heading("Описание", level=2)
p = doc.add_paragraph(
    "Классический невысокий военный городок с двух- и трёхэтажными домами, дворами и узкими проходами между зданиями. "
    "Локация строится вокруг свободного выбора маршрута: игрок может идти по земле, забираться на крыши "
    "или проходить дома насквозь."
)
p.paragraph_format.space_after = Pt(8)

doc.add_heading("Новая механика", level=2)
p = doc.add_paragraph(
    "На этапе игрок знакомится с прыжками в окна. Прыжок позволяет быстро менять маршрут, врываться внутрь зданий "
    "и продолжать движение через комнаты к противоположному выходу."
)
p.paragraph_format.space_after = Pt(8)

doc.add_heading("Бой и маршруты", level=2)
p = doc.add_paragraph(
    "Часть противников патрулирует крыши, часть остаётся на земле. Игрок сам выбирает высоту и путь: "
    "сражаться на улицах, обходить врагов по крышам или использовать здания для укрытия и флангового прохода."
)
p.paragraph_format.space_after = Pt(12)

doc.add_heading("6 Дорога Бой с БТР", level=1)

summary = doc.add_table(rows=3, cols=2)
summary.style = "Table Grid"
summary.autofit = False
rows = [
    ("Статус", "Обязательный бой с техникой"),
    ("Угроза", "БТР разрушает почти все стационарные укрытия"),
    ("Задача", "Уклоняться от огня, быстро менять позицию и уничтожить БТР гранатами"),
]
for row, (label, value) in zip(summary.rows, rows):
    row.cells[0].width = Inches(1.4)
    row.cells[1].width = Inches(5.0)
    set_cell_shading(row.cells[0], "E7EBEF")
    p = row.cells[0].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(label), size=9.5, bold=True)
    p = row.cells[1].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(value), size=9.5)

doc.add_paragraph()
doc.add_heading("Описание", level=2)
p = doc.add_paragraph(
    "Дорога от военного городка выводит игрока на бой с БТР. Машина атакует издалека и методично "
    "разрушает почти все стационарные укрытия, поэтому нельзя переждать обстрел в одном месте."
)
p.paragraph_format.space_after = Pt(8)

doc.add_heading("Тактика и механики", level=2)
p = doc.add_paragraph(
    "Игроку нужно постоянно бегать и менять укрытия. На дороге стоят ящики с гранатами: ими требуется "
    "закидывать БТР. Для коротких перебежек герой вырывает двери припаркованных машин и использует их "
    "как переносную защиту."
)
p.paragraph_format.space_after = Pt(12)

doc.add_heading("Кадр появления", level=2)
add_btr_intro(doc)

doc.add_heading("7 Военная часть", level=1)

summary = doc.add_table(rows=3, cols=2)
summary.style = "Table Grid"
summary.autofit = False
rows = [
    ("Статус", "Обязательная вертикальная стелс секция"),
    ("Цель", "Проникнуть внутрь высокой военной части и выйти на крышу"),
    ("Ограничение", "На фасад нельзя запрыгнуть, а все окна защищены решётками"),
]
for row, (label, value) in zip(summary.rows, rows):
    row.cells[0].width = Inches(1.4)
    row.cells[1].width = Inches(5.0)
    set_cell_shading(row.cells[0], "E7EBEF")
    p = row.cells[0].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(label), size=9.5, bold=True)
    p = row.cells[1].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(value), size=9.5)

doc.add_paragraph()
doc.add_heading("Фасад и вход", level=2)
p = doc.add_paragraph(
    "Военная часть представляет собой высокое П образное здание с главным входом на фасаде. "
    "Его высота исключает прыжок на крышу, а решётки на окнах закрывают привычный путь через оконные проёмы. "
    "Игроку нужно попасть внутрь через главный вход или другой наземный проход."
)
p.paragraph_format.space_after = Pt(8)

doc.add_heading("Подъём на крышу", level=2)
p = doc.add_paragraph(
    "После проникновения начинается маршрут вверх: герой проходит через внутренние помещения, "
    "находит лестничный путь и поднимается на крышу. Крыша становится точкой выхода к следующему этапу."
)
p.paragraph_format.space_after = Pt(8)

doc.add_heading("Интерьер и стелс", level=2)
p = doc.add_paragraph(
    "Внутри военной части коридоры разбиты бетонными стенками и перегородками. Игрок скрывается за ними, "
    "наблюдает за патрулями и продвигается от укрытия к укрытию. Открытый бой возможен, но планировка прежде всего "
    "поддерживает осторожное скрытное прохождение."
)
p.paragraph_format.space_after = Pt(12)

doc.add_heading("Кадры сцены", level=2)
add_military_unit_storyboard(doc)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Wolfish Revenge  Рабочий документ")
set_run_font(footer_run, size=8.5, color=(120, 120, 120))

doc.save(OUTPUT)
print(OUTPUT)
